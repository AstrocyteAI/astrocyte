"""Unit tests for astrocyte_ingestion_document._extract — no I/O beyond tmp files."""

from __future__ import annotations

import io

import pytest

from astrocyte_ingestion_document._extract import extract_text

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _write(tmp_path, name: str, content: bytes) -> "Path":  # noqa: F821
    from pathlib import Path

    p: Path = tmp_path / name
    p.write_bytes(content)
    return p


# ---------------------------------------------------------------------------
# Plain text
# ---------------------------------------------------------------------------


def test_plain_txt(tmp_path) -> None:
    p = _write(tmp_path, "notes.txt", "Hello, world!\nLine two.".encode("utf-8"))
    result = extract_text(p)
    assert result is not None
    assert "Hello, world!" in result


def test_plain_md(tmp_path) -> None:
    p = _write(tmp_path, "README.md", "# Title\n\nSome content.".encode("utf-8"))
    result = extract_text(p)
    assert result is not None
    assert "# Title" in result


def test_plain_latin1_fallback(tmp_path) -> None:
    p = _write(tmp_path, "legacy.txt", "caf\xe9".encode("latin-1"))
    result = extract_text(p)
    assert result is not None
    assert "caf" in result


def test_plain_csv(tmp_path) -> None:
    p = _write(tmp_path, "data.csv", b"col1,col2\nval1,val2\n")
    result = extract_text(p)
    assert result is not None
    assert "col1" in result


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------


def test_html_strips_script(tmp_path) -> None:
    body = b"<html><body><script>evil()</script><p>Keep.</p></body></html>"
    p = _write(tmp_path, "page.html", body)
    result = extract_text(p)
    assert result is not None
    assert "Keep." in result
    assert "evil" not in result


def test_htm_extension(tmp_path) -> None:
    p = _write(tmp_path, "index.htm", b"<p>Old school.</p>")
    result = extract_text(p)
    assert result is not None
    assert "Old school." in result


# ---------------------------------------------------------------------------
# PDF (requires pypdf)
# ---------------------------------------------------------------------------


def _make_minimal_pdf(text: str) -> bytes:
    try:
        from fpdf import FPDF  # type: ignore[import-untyped]

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(200, 10, txt=text)
        return pdf.output()
    except ImportError:
        pass

    stream = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode()
    slen = len(stream)
    return (
        b"%PDF-1.4\n"
        b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
        b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]"
        b"/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj\n"
        + f"4 0 obj<</Length {slen}>>stream\n".encode()
        + stream
        + b"\nendstream endobj\n"
        b"5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\n"
        b"xref\n0 6\n"
        b"0000000000 65535 f \n"
        b"0000000009 00000 n \n"
        b"0000000058 00000 n \n"
        b"0000000115 00000 n \n"
        b"0000000266 00000 n \n"
        b"0000000400 00000 n \n"
        b"trailer<</Size 6/Root 1 0 R>>\n"
        b"startxref\n460\n%%EOF"
    )


def test_pdf_no_crash(tmp_path) -> None:
    pytest.importorskip("pypdf")
    p = _write(tmp_path, "doc.pdf", _make_minimal_pdf("PDF content"))
    result = extract_text(p)
    assert result is None or isinstance(result, str)


def test_pdf_corrupt_returns_none(tmp_path) -> None:
    pytest.importorskip("pypdf")
    p = _write(tmp_path, "bad.pdf", b"not a real pdf")
    result = extract_text(p)
    assert result is None


# ---------------------------------------------------------------------------
# DOCX (requires python-docx)
# ---------------------------------------------------------------------------


def _make_docx(text: str) -> bytes:
    from docx import Document  # type: ignore[import-untyped]

    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_extraction(tmp_path) -> None:
    pytest.importorskip("docx")
    p = _write(tmp_path, "report.docx", _make_docx("Hello from DOCX"))
    result = extract_text(p)
    assert result is not None
    assert "Hello from DOCX" in result


def test_docx_corrupt_returns_none(tmp_path) -> None:
    pytest.importorskip("docx")
    p = _write(tmp_path, "bad.docx", b"not a real docx")
    result = extract_text(p)
    assert result is None


# ---------------------------------------------------------------------------
# Size cap
# ---------------------------------------------------------------------------


def test_oversized_file_returns_none(tmp_path, monkeypatch) -> None:
    from astrocyte_ingestion_document import _extract as mod

    monkeypatch.setattr(mod, "MAX_BYTES", 10)
    p = _write(tmp_path, "big.txt", b"x" * 11)
    result = extract_text(p)
    assert result is None
